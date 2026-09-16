"""Read-only imports of the two approved cue sheets; no Google write API."""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime, time
from html import escape
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SOURCES = {
    "주일예배": "1baPrJ5Tg12g8SPK8FTz6d_RnGn-MIX6T3YWxcwg-Hqc",
    "금요집회": "1L1401SkPbj3FBb0YwJXA7pTnx31QVVXIydzKX1vyhOo",
}
DATE = re.compile(r"(20\d{2})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일")


def norm(value):
    return re.sub(r"\s+", "", str(value or "")).casefold()


def key(value):
    return hashlib.sha256(value.encode()).hexdigest()[:24]


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, time)):
        return value.strftime("%H:%M")
    return str(value).strip()


def download(kind):
    if kind not in SOURCES:
        raise ValueError("등록되지 않은 큐시트입니다.")
    request = Request(f"https://docs.google.com/spreadsheets/d/{SOURCES[kind]}/export?format=xlsx",
                      headers={"User-Agent": "Joyful-Cue-ReadOnly/1.0"})
    try:
        with urlopen(request, timeout=25) as response:
            data = response.read(20 * 1024 * 1024 + 1)
        if len(data) > 20 * 1024 * 1024:
            raise ValueError()
        return parse_workbook(data, kind)
    except Exception:
        raise ValueError("큐시트를 읽지 못했어요. 공유 권한·인터넷 연결을 확인하세요. 기존 확정본은 유지됩니다.") from None


def parse_workbook(data, kind):
    if kind not in SOURCES:
        raise ValueError("등록되지 않은 큐시트입니다.")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError("큐시트 크기가 너무 큽니다.")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    plans = []
    try:
        if len(workbook.sheetnames) > 40:
            raise ValueError("탭이 너무 많아요. 범위를 확인해 주세요.")
        for sheet in workbook:
            rows = [[text(c) for c in row] for row in sheet.iter_rows(max_row=300, max_col=60, values_only=True)]
            headers = [(r, c, DATE.search(v)) for r, row in enumerate(rows) for c, v in enumerate(row)
                       if DATE.search(v) and (("주일예배" in norm(v)) if kind == "주일예배" else ("금요" in v))]
            for r, c, match in headers:
                try:
                    day = date(*map(int, match.groups())).isoformat()
                except ValueError:
                    continue
                right = min([cc for rr, cc, _ in headers if rr == r and cc > c] or [min(c+15, 60)])
                bottom = min([rr for rr, cc, _ in headers if cc == c and rr > r] or [len(rows)])
                header_row = next((rr for rr in range(r+1, min(r+18, bottom))
                                   if "내용" in rows[rr][c:right] and "담당" in rows[rr][c:right]), None)
                if header_row is None:
                    continue
                columns = {norm(v): cc for cc, v in enumerate(rows[header_row]) if c <= cc < right and v}
                def get(rr, label):
                    col = columns.get(norm(label))
                    return rows[rr][col] if col is not None else ""
                items = []
                for rr in range(header_row+1, bottom):
                    body = "\n".join(v for v in rows[rr][columns["내용"]:columns["담당"]]
                                     if v and not re.fullmatch(r"[\d\s]+", v)).strip()
                    if not body or body == "`":
                        continue
                    owner = get(rr, "담당")
                    base = dict(time=get(rr, "시간"), duration=get(rr, "R . T"), owner=owner,
                                sound=get(rr, "음향"), screen=get(rr, "화면"), light=get(rr, "조명"),
                                stage=get(rr, "무대"), notes=get(rr, "세부사항"),
                                cell=f"{get_column_letter(columns['내용']+1)}{rr+1}")
                    try:
                        light = float(base["light"])
                        if 0 <= light <= 1:
                            base["light"] = f"{light*100:g}%"
                    except ValueError:
                        pass
                    # Continuation rows belong to their cue, not a new clock event.
                    if not owner and items and items[-1]["kind"] in {"본문", "설교 후 찬양", "기도회"}:
                        items[-1]["body"] += ("\n" if items[-1]["body"] else "") + body
                        continue
                    if norm(body) == "성경봉독":
                        category, title, content = "본문", "성경봉독", ""
                    elif norm(body) == "설교후찬양":
                        category, title, content = "설교 후 찬양", "설교 후 찬양", ""
                    elif norm(body) == "기도회":
                        category, title, content = "기도회", "기도회", ""
                    elif norm(owner) == "뿌리깊은나무" and "축도" not in body:
                        for song in body.splitlines():
                            if song.strip():
                                items.append(dict(base, kind="뿌나 찬양", title=song.strip(), body=song.strip()))
                        continue
                    else:
                        category, title, content = "순서", body, body
                    items.append(dict(base, kind=category, title=title, body=content))
                if not items:
                    continue
                occurrences = {}
                for item in items:
                    semantic = item["kind"]+":"+norm(item["title"])
                    occurrences[semantic] = occurrences.get(semantic, 0)+1
                    item["id"] = key(semantic+str(occurrences[semantic]))
                    if item["kind"] == "설교 후 찬양" and item["body"]:
                        item["title"] = "설교 후 찬양 · " + item["body"]
                checks = []
                check_col = next((cc for rr in range(r+1, header_row) for cc in range(c, right)
                                  if norm(rows[rr][cc]) == "체크사항"), None)
                if check_col is not None:
                    for rr in range(r+1, header_row):
                        for line in rows[rr][check_col].splitlines():
                            line = line.strip(" -\t")
                            if line and norm(line) != "체크사항" and line not in checks:
                                checks.append(line)
                warnings = ["찬양 키·악보 링크는 자동 추측하지 않습니다. 연결할 자료는 별도 확인하세요."]
                if len([1 for rr, _, _ in headers if rr == r]) > 1:
                    warnings.append("같은 탭에 여러 날짜가 있어요. 선택한 날짜·영역만 가져왔습니다.")
                if not any(i["kind"] == "본문" and i["body"] for i in items):
                    warnings.append("본문 말씀을 찾지 못했어요. 원본을 확인하세요.")
                if not any(i["kind"] == "뿌나 찬양" for i in items):
                    warnings.append("뿌나 찬양을 찾지 못했어요. 원본을 확인하세요.")
                if any(i["kind"] == "설교 후 찬양" and not i["body"] for i in items):
                    warnings.append("설교 후 찬양 제목이 비어 있어요. 원본을 확인하세요.")
                warnings.append("지원 범위: 탭별 앞 300행·60열. 미리보기와 원본을 확인한 뒤 사용하세요.")
                source_id = SOURCES[kind]
                plan = dict(id=key(source_id+sheet.title+day+str(c)), kind=kind, date=day,
                            sheet=sheet.title, area=f"{get_column_letter(c+1)}{r+1}:{get_column_letter(right)}{bottom}",
                            source_id=source_id, url=f"https://docs.google.com/spreadsheets/d/{source_id}/edit",
                            title=f"{day} {kind}", items=items, checks=checks, warnings=warnings)
                plan["hash"] = key(json.dumps(plan, ensure_ascii=False, sort_keys=True))
                plans.append(plan)
    finally:
        workbook.close()
    if not plans:
        raise ValueError("지원하는 날짜·내용·담당 표를 찾지 못했어요. 원본 구조를 확인하세요.")
    return sorted(plans, key=lambda p: p["date"], reverse=True)


def differences(old, new):
    a = {i["id"]: i for i in old["items"]}
    b = {i["id"]: i for i in new["items"]}
    lines = ["추가: "+b[k]["title"] for k in b.keys()-a.keys()]
    lines += ["삭제: "+a[k]["title"] for k in a.keys()-b.keys()]
    lines += ["변경: "+b[k]["title"] for k in a.keys() & b.keys() if a[k] != b[k]]
    if list(a) != list(b):
        lines.append("순서 또는 목록 구성이 변경됐어요.")
    if old["checks"] != new["checks"]:
        lines.append("준비 체크 항목이 변경됐어요.")
    return lines


def offline_html(plan):
    chunks = ["<!doctype html><html lang='ko'><meta charset='utf-8'><meta name='viewport' content='width=device-width'><title>예배 보기용 사본</title><style>body{font:20px sans-serif;max-width:900px;margin:30px auto;padding:20px;line-height:1.6}article{border-top:1px solid #aaa;padding:15px 0}pre{white-space:pre-wrap}button{font-size:20px}@media print{button{display:none}}</style>",
              "<h1>"+escape(plan["title"])+"</h1><p>보기용 사본 · 실시간 요청·진행 동기화는 작동하지 않습니다. 내부 팀 자료입니다.</p><button onclick='window.print()'>인쇄</button>"]
    for item in plan["items"]:
        chunks.append("<article><h2>"+escape(item["title"])+"</h2><pre>"+escape(item["body"])+"</pre>")
        for label, field in (("음향", "sound"), ("화면", "screen"), ("조명", "light"), ("무대", "stage"), ("준비", "notes")):
            if item[field]:
                chunks.append("<p>"+label+": "+escape(item[field])+"</p>")
        chunks.append("</article>")
    return "".join(chunks)+"</html>"
